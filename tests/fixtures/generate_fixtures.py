"""生成 docx→Markdown 转换器的测试夹具（可重复运行）。

用法（Python 需装 python-docx 与 Pillow）：
  python tests/fixtures/generate_fixtures.py

产物提交到仓库，Node 测试直接读取，不需要 Python 运行时。
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Pt, Inches
from PIL import Image, ImageDraw

HERE = Path(__file__).resolve().parent
OUT = HERE / "docx"
OUT.mkdir(parents=True, exist_ok=True)


def add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def make_png(color: tuple[int, int, int], label: str) -> bytes:
    img = Image.new("RGB", (240, 120), color)
    draw = ImageDraw.Draw(img)
    draw.text((16, 52), label, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def add_monospace_paragraph(doc, lines: list[str]):
    para = doc.add_paragraph()
    for i, line in enumerate(lines):
        if i > 0:
            br_run = para.add_run()
            br_run.font.name = "Consolas"
            br_run.font.size = Pt(10)
            br_run.add_break()
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
    return para


def set_repeat_table_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_basic() -> None:
    doc = Document()

    title = doc.add_paragraph("DocMax 导入测试", style="Title")
    if title.style is None:  # pragma: no cover - default 模板总有 Title
        title.style = "Heading 1"

    doc.add_heading("第一章 基本结构", level=1)

    p = doc.add_paragraph()
    p.add_run("这是一段包含 ")
    p.add_run("粗体文字").bold = True
    p.add_run("、")
    p.add_run("斜体文字").italic = True
    p.add_run(" 和普通文字的中文段落，用于验证行内格式。")

    doc.add_paragraph("无序列表第一项", style="List Bullet")
    doc.add_paragraph("无序列表第二项（含粗体）", style="List Bullet")
    doc.add_paragraph("有序列表第一项", style="List Number")
    doc.add_paragraph("有序列表第二项", style="List Number")

    doc.add_heading("1.1 小节标题", level=2)
    link_para = doc.add_paragraph()
    link_para.add_run("参考链接：")
    add_hyperlink(link_para, "https://api-docs.deepseek.com/", "DeepSeek API 文档")
    link_para.add_run(" 以及后续说明。")

    # 没有标记表头行的表格：转换器应生成空表头并保留所有数据行
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    header = ["姓名", "角色", "备注"]
    for i, text in enumerate(header):
        table.cell(0, i).text = text
    data = [["张三", "产品", "负责需求"], ["李四", "工程", "负责实现"], ["王五", "设计", "负责视觉"]]
    for r, row in enumerate(data, start=1):
        for c, text in enumerate(row):
            table.cell(r, c).text = text

    doc.add_heading("第二章 代码与引用", level=1)

    # Word 里代码通常是多行：每行一个段落（代码样式）+ 一个段内软换行，两种都测
    code_para = doc.add_paragraph()
    run = code_para.add_run("const answer: number = 42;")
    run.font.name = "Consolas"
    run.font.size = Pt(10)

    add_monospace_paragraph(doc, ["function add(a: number, b: number) {", "  return a + b;", "}"])

    quote = doc.add_paragraph("引用块里的观点：编辑工具不应该打断写作。")
    try:
        quote.style = "Quote"
    except KeyError:  # pragma: no cover
        pass

    doc.add_paragraph("最后一段普通正文。")
    doc.save(OUT / "basic.docx")


def build_table_header() -> None:
    doc = Document()
    doc.add_heading("表头测试", level=1)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "列 A"
    table.cell(0, 1).text = "列 B"
    table.cell(1, 0).text = "a1"
    table.cell(1, 1).text = "b1"
    table.cell(2, 0).text = "a2"
    table.cell(2, 1).text = "b2"
    set_repeat_table_header(table.rows[0])
    doc.save(OUT / "table-header.docx")


def build_images() -> None:
    doc = Document()
    doc.add_heading("图文混排测试", level=1)
    doc.add_paragraph("下面是一张示意图：")
    doc.add_picture(io.BytesIO(make_png((37, 99, 235), "diagram")), width=Inches(3))
    doc.add_paragraph("图片说明：这是图 1。")
    doc.add_picture(io.BytesIO(make_png((194, 65, 12), "screenshot")), width=Inches(3))
    doc.add_paragraph("最后一段正文，验证图片前后的文字顺序。")
    doc.save(OUT / "images.docx")


def build_empty() -> None:
    doc = Document()
    doc.save(OUT / "empty.docx")


def build_headings_only() -> None:
    doc = Document()
    doc.add_heading("一级标题 A", level=1)
    doc.add_paragraph("正文一。")
    doc.add_heading("二级标题 A.1", level=2)
    doc.add_paragraph("正文二。")
    doc.add_heading("三级标题 A.1.1", level=3)
    doc.add_paragraph("正文三。")
    doc.save(OUT / "headings.docx")


def build_unsafe_link() -> None:
    doc = Document()
    doc.add_heading("危险链接测试", level=1)
    p = doc.add_paragraph("不要点击：")
    add_hyperlink(p, "javascript:alert(1)", "恶意链接")
    p2 = doc.add_paragraph("正常链接：")
    add_hyperlink(p2, "https://example.com/", "示例站点")
    doc.save(OUT / "unsafe-link.docx")


if __name__ == "__main__":
    build_basic()
    build_images()
    build_empty()
    build_headings_only()
    build_table_header()
    build_unsafe_link()
    for path in sorted(OUT.glob("*.docx")):
        print(f"{path.name}: {path.stat().st_size} bytes")
